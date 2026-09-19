import logging
from pathlib import Path
from typing import Union

import docx
import pypdf

logger = logging.getLogger(__name__)


def extract_text_from_txt(file_path: Path) -> str:
    """Read plain text file with encoding fallback."""
    encodings = ["utf-8", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                text = f.read().strip()
                if text:
                    return text
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Could not decode text file {file_path.name} with standard encodings.")


def extract_text_from_pdf(file_path: Path) -> str:
    """Extract text content from a PDF document using pypdf."""
    text_parts = []
    try:
        reader = pypdf.PdfReader(file_path)
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text.strip())
    except Exception as e:
        logger.error(f"Error parsing PDF {file_path.name}: {e}")
        raise ValueError(f"Failed to parse PDF document {file_path.name}: {str(e)}")

    full_text = "\n\n".join(text_parts).strip()
    if not full_text:
        raise ValueError(f"PDF document {file_path.name} contains no readable text.")
    return full_text


def extract_text_from_docx(file_path: Path) -> str:
    """Extract text content from a DOCX document including paragraphs and tables."""
    text_parts = []
    try:
        doc = docx.Document(file_path)
        # Extract paragraph text
        for p in doc.paragraphs:
            if p.text.strip():
                text_parts.append(p.text.strip())
        # Extract text from tables if present
        for table in doc.tables:
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_data:
                    text_parts.append(" | ".join(row_data))
    except Exception as e:
        logger.error(f"Error parsing DOCX {file_path.name}: {e}")
        raise ValueError(f"Failed to parse DOCX document {file_path.name}: {str(e)}")

    full_text = "\n".join(text_parts).strip()
    if not full_text:
        raise ValueError(f"DOCX document {file_path.name} contains no readable text.")
    return full_text


def extract_text_from_file(file_path: Union[str, Path]) -> str:
    """
    Main entry point for extracting text from a file.
    Supports .txt, .pdf, and .docx formats.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".txt":
        return extract_text_from_txt(path)
    elif suffix == ".pdf":
        return extract_text_from_pdf(path)
    elif suffix == ".docx":
        return extract_text_from_docx(path)
    else:
        raise ValueError(f"Unsupported file format '{suffix}' for file {path.name}. Supported formats: .txt, .pdf, .docx")
